"""Build the project handoff report as a polished DOCX."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/CallMetricAI_Stereo_Whisper_Teknik_Degerlendirme_Raporu.docx")
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 78, 121)
MUTED = RGBColor(90, 99, 110)
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F4F7"
GREEN = "E8F5E9"
AMBER = "FFF4CE"
RED = "FDECEC"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def format_table(table, widths: list[int], header=True) -> None:
    set_table_geometry(table, widths)
    if header:
        set_repeat_table_header(table.rows[0])
        for cell in table.rows[0].cells:
            set_cell_shading(cell, LIGHT_BLUE)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = DARK_BLUE
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(9)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, text in enumerate(headers):
        table.rows[0].cells[i].text = text
    for values in rows:
        cells = table.add_row().cells
        for i, text in enumerate(values):
            cells[i].text = text
    format_table(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(text, style=style)
    p.paragraph_format.space_after = Pt(3)


def add_number(doc, text: str) -> None:
    p = doc.add_paragraph(text, style="List Number")
    p.paragraph_format.space_after = Pt(4)


def add_callout(doc, label: str, text: str, fill: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    r = p.add_run(f"{label}: ")
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    p.add_run(text)
    set_table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Bullet 2", "List Number"):
        styles[name].font.name = "Calibri"
        styles[name].font.size = Pt(10.5)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Sayfa ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def build() -> Path:
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    header = section.header.paragraphs[0]
    header.text = "CallMetricAI  |  Stereo Whisper Teknik Değerlendirme"
    header.style = doc.styles["Normal"]
    header.runs[0].font.size = Pt(8.5)
    header.runs[0].font.color.rgb = MUTED
    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = MUTED

    # Cover / memo masthead
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("TEKNİK DEĞERLENDİRME RAPORU")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = BLUE
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("CallMetricAI Offline Stereo Whisper İş Hattı")
    r.bold = True
    r.font.size = Pt(25)
    r.font.color.rgb = DARK_BLUE
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run(
        "Modüler mimari, zaman hizalama, gerçek veri doğrulaması, "
        "modüler mimari, denetim zinciri, gerçek veri doğrulaması ve teknik sınırlamalar"
    )
    r.font.size = Pt(13)
    r.font.color.rgb = MUTED

    metadata = [
        ("Proje", "simple_console_whisper_ollama"),
        ("Durum", "Offline stereo kabul zinciri doğrulandı; otomatik doğrulama PASS"),
        ("Kapsam", "Dosyadan stereo işleme; canlı entegrasyon kapsam dışı"),
        ("Rapor tarihi", date.today().strftime("%d.%m.%Y")),
        ("Test veri seti", "HarperValleyBank SID 0002f70f7386445b"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{label}: ")
        r.bold = True
        r.font.color.rgb = DARK_BLUE
        p.add_run(value)

    add_callout(
        doc,
        "Teknik sonuç",
        "Mevcut main.py değiştirilmeden stereo iş hattı yan modüller olarak "
        "eklenmiştir. Dokuz otomatik test geçmekte; gerçek çağrıda beklenen 2 "
        "çakışmanın 2'si de bulunmakta, örnek ekseninde ve ASR zamanlarında drift "
        "tespit edilmemekte, tüm aşamalar audit.json içinde izlenmektedir.",
        GREEN,
    )

    doc.add_heading("1. Yönetici özeti", level=1)
    doc.add_paragraph(
        "Çalışmanın amacı, mevcut mono canlı uygulamayı değiştirmeden yanına "
        "offline stereo görüşmeleri işleyen, sınıf tabanlı ve izlenebilir bir iş "
        "hattı eklemektir. Temel tasarım kararı stereo dosyayı doğrudan Whisper'a "
        "vermemek; sol ve sağ kanalı adlandırılmış iki mono akış olarak ayrı ayrı "
        "işlemek, sonuçları ortak zaman ekseninde birleştirmek ve yalnız final "
        "timeline'ı analiz katmanına göndermektir."
    )
    add_bullet(doc, "Mevcut main.py canlı sistemi değiştirilmemiştir.")
    add_bullet(doc, "Sol kanal agent, sağ kanal customer olarak sözleşmeye bağlanmıştır.")
    add_bullet(doc, "8/16/48 kHz girdiler 16 kHz mono kanonik formata dönüştürülebilir.")
    add_bullet(doc, "Faster-Whisper ve Ollama arayüzlerin arkasında değiştirilebilir adaptörlerdir.")
    add_bullet(doc, "Her aşama JSON ve Markdown çıktıda süre ve sonuç olarak izlenmektedir.")
    add_bullet(doc, "Dokuz otomatik testin tamamı geçmiştir; gerçek görüşmede uçtan uca zincir PASS üretmiştir.")
    add_bullet(doc, "Referans dosyası yalnız doğrulama oracle'ıdır; transkripti veya timeline'ı düzeltmez.")

    doc.add_heading("2. Kapsam ve kabul kriterleri", level=1)
    add_table(
        doc,
        ["Kriter", "Durum", "Kanıt / değerlendirme"],
        [
            ["Mevcut sistemi bozmama", "Karşılandı", "main.py üzerinde değişiklik yok."],
            ["Sınıf tabanlı modüler mimari", "Karşılandı", "Audio, ASR, timeline, analiz, orkestrasyon ve raporlama ayrıldı."],
            ["İki kanalı doğru ayırma", "Karşılandı", "L=agent, R=customer; PCM16 ve kanal sayısı doğrulanıyor."],
            ["Zaman eksenini koruma", "Karşılandı", "Sessizlikler fiziksel olarak kesilip uç uca eklenmiyor."],
            ["Doğru transkripsiyon", "Kısmen", "Caller WER %3,03; agent WER %14,89 (tiny model, 8 kHz kaynak)."],
            ["Doğru timeline/overlap", "Karşılandı", "Word-gap split + kanal VAD; beklenen 2, bulunan 2."],
            ["Analize gönderme", "Kısmen", "İstek/yanıt eksiksiz loglanıyor; conversation_stage semantik olarak hatalı."],
            ["İzlenebilir test ve rapor", "Karşılandı", "9 test, audit.json, result.json, Markdown ve bu teknik rapor."],
            ["Canlı entegrasyon", "Kapsam dışı", "Bu fazda özellikle uygulanmadı."],
        ],
        [2500, 1350, 5510],
    )

    doc.add_heading("3. Mimari ve veri akışı", level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Stereo WAV → Doğrulama/L-R ayırma → HQ resample → Kanal VAD + Whisper → "
        "Word-gap split → Zaman sıralı merge → Ollama → Audit/validasyon/rapor"
    )
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    r.font.size = Pt(11)

    add_table(
        doc,
        ["Modül", "Sınıf / arayüz", "Sorumluluk"],
        [
            ["audio.py", "StereoAudioProcessor", "Girdi doğrulama, kanal ayırma, SOXR HQ resample, kalite ölçümü."],
            ["models.py", "Domain dataclass'ları", "Kanal, kelime, segment, timeline, kalite ve işlem sonucu modelleri."],
            ["transcription.py", "ChannelTranscriber", "ASR bağımlılığından bağımsız kanal transkripsiyon sözleşmesi."],
            ["transcription.py", "FasterWhisperTranscriber", "Tek model örneğiyle iki kanalı sıralı ve ayrı işleme."],
            ["activity.py", "EnergySpeechActivityDetector", "Her kanalda bağımsız konuşma aralığı ve gerçek overlap otoritesi."],
            ["postprocessing.py", "TranscriptSegmentSplitter", "Kelime zamanlarındaki sessizlik boşluklarından otomatik segment bölme."],
            ["timeline.py", "TimelineMerger", "Kararlı kronolojik sıralama ve VAD tabanlı kanal çakışması."],
            ["analysis.py", "ConversationAnalyzer", "Analiz backend'inden bağımsız sözleşme."],
            ["analysis.py", "OllamaConversationAnalyzer", "JSON çıktı, confidence ve evidence ID doğrulaması."],
            ["pipeline.py", "StereoProcessingPipeline", "Aşama orkestrasyonu ve süre izleri."],
            ["validation.py", "ReferenceTranscriptValidator", "Kanal WER, kelime zamanı drift'i ve overlap golden doğrulaması."],
            ["reporting.py", "ResultReporter", "result.json, sıralı audit.json ve insan okunur Markdown rapor."],
            ["cli.py", "CLI", "Offline dosya işleme giriş noktası ve dil/model yapılandırması."],
        ],
        [1800, 2700, 4860],
    )

    doc.add_heading("4. Zaman hizalama ve kanal bütünlüğü", level=1)
    doc.add_paragraph(
        "Stereo dosya tek iki boyutlu örnek dizisi olarak okunur. Her frame'in "
        "0. kanalı agent, 1. kanalı customer olarak çıkarılır. Kanalların başındaki "
        "ve arasındaki sessizlikler korunur. Kaynak örnekleme hızı 16 kHz değilse "
        "her iki kanala aynı SOXR HQ dönüşümü uygulanır. Olası tek örneklik "
        "yuvarlama farkı ortak minimum uzunlukla simetrik biçimde giderilir."
    )
    add_callout(
        doc,
        "Önemli ayrım",
        "Pipeline kaynaklı zaman kayması ile ASR timestamp belirsizliği aynı şey "
        "değildir. Örnek ekseni korunurken Whisper kelime/segment sınırları gerçek "
        "enerji başlangıcından birkaç yüz milisaniye önce veya sonra tahmin edebilir.",
        LIGHT_BLUE,
    )
    add_bullet(doc, "Kanal bazında bağımsız sessizlik kırpma yapılmaz.")
    add_bullet(doc, "Kalan konuşma parçaları uç uca eklenmez.")
    add_bullet(doc, "Aynı anda konuşmalar tek metne ezilmez; iki segment de korunur.")
    add_bullet(doc, "Dual-mono şüphesi için kanal korelasyonu raporlanır.")
    add_bullet(doc, "Birleşik ölçümde medyan zaman farkı 70 ms, p95 mutlak hata 306 ms ve eğim 127,38 ms/dk'dır.")
    add_bullet(doc, "Tanımlı eşikler 250 ms/dk ve p95 500 ms olduğundan ASR timestamp drift'i yoktur.")

    doc.add_heading("5. Test yaklaşımı", level=1)
    doc.add_heading("5.1 Deterministik otomatik testler", level=2)
    add_table(
        doc,
        ["Test", "Sonuç"],
        [
            ["Bağımsız kanal VAD ile gerçek konuşma/overlap tespiti", "Geçti"],
            ["16 kHz L/R kanal kimliği ve 48.000 örneklik ortak süre", "Geçti"],
            ["48 kHz → 16 kHz HQ resample ve eşit kanal uzunluğu", "Geçti"],
            ["Mono girdinin stereo olarak reddedilmesi", "Geçti"],
            ["Kanal bazlı RMS, sessizlik ve korelasyon ölçümü", "Geçti"],
            ["Kararlı timeline sıralaması ve overlap korunması", "Geçti"],
            ["Uçtan uca orkestrasyon, analiz ve iz kayıtları", "Geçti"],
            ["Kelime boşluğundan uzun Whisper segmentini otomatik bölme", "Geçti"],
            ["Referans WER, drift ve overlap validasyonu", "Geçti"],
        ],
        [7000, 2360],
    )
    doc.add_paragraph("Test komutu:  .\\venv\\Scripts\\python.exe -m unittest discover -s tests -v")

    doc.add_heading("5.2 Kontrollü sentetik stereo testi", level=2)
    doc.add_paragraph(
        "Windows konuşma senteziyle oluşturulan iki mono örnek ortak zaman eksenine "
        "yerleştirilmiş; agent ve customer metinleri doğru kanalda çevrilmiş, bir "
        "çakışma korunmuş ve Ollama bütçe/fiyat itirazını customer segmentine "
        "dayandırmıştır. Fiziksel kanal başlangıç farkı 1,002 sn; hedef 1,000 sn'dir."
    )

    doc.add_heading("5.3 Gerçek HarperValleyBank görüşmesi", level=2)
    add_table(
        doc,
        ["Alan", "Değer"],
        [
            ["SID", "0002f70f7386445b"],
            ["Kaynak", "HarperValleyBank; insanlarla canlandırılmış banka çağrısı"],
            ["Lisans", "CC BY 4.0"],
            ["Kaynak ses", "Agent ve caller için ayrı 8 kHz PCM16 mono WAV"],
            ["Üretilen fixture", "16 kHz PCM16 stereo; L=agent, R=customer"],
            ["Süre", "51,110 sn"],
            ["Referans", "18 insan etiketli segment; speaker_role ve zaman damgaları"],
        ],
        [2600, 6760],
    )

    doc.add_heading("6. Gerçek veri sonuçları", level=1)
    doc.add_heading("6.1 Ses ve kanal kalitesi", level=2)
    add_table(
        doc,
        ["Metrik", "Agent", "Customer", "Yorum"],
        [
            ["RMS", "-36,84 dBFS", "-30,99 dBFS", "Agent yaklaşık 5,85 dB daha düşük seviyede."],
            ["Clipping", "%0,000", "%0,000", "Kırpılma yok."],
            ["Sessizlik", "%86,16", "%86,04", "Konuşma seyrek; kanal ekseni sessizliklerle korunuyor."],
            ["Korelasyon", "0,000154", "0,000154", "Dual-mono değil; kanallar bağımsız."],
        ],
        [1900, 1500, 1500, 4460],
    )

    doc.add_heading("6.2 ASR doğruluğu", level=2)
    add_table(
        doc,
        ["Kanal", "Referans kelime", "Whisper kelime", "Hata", "WER"],
        [
            ["Agent", "47", "48", "7", "%14,89"],
            ["Customer", "33", "32", "1", "%3,03"],
        ],
        [2000, 1900, 1900, 1400, 2160],
    )
    doc.add_paragraph(
        "Caller/customer transkripsiyonu güçlüdür. Agent tarafındaki hata ağırlıklı "
        "olarak düşük seviyeli 8 kHz telefon sesinde banka karşılama cümlesi ve "
        "\"Which card would you like to replace?\" ifadesinde görülmüştür. "
        "Bu sonuç tiny/cpu/int8 konfigürasyonuna aittir; small model benchmark'ı "
        "kapanış öncesi önerilmektedir. Customer kanalındaki tek hata, kısa "
        "\"Hi\" kelimesinin atlanmasıdır."
    )

    doc.add_heading("6.3 Performans", level=2)
    add_table(
        doc,
        ["Aşama", "Süre", "Değerlendirme"],
        [
            ["Yükleme ve kanal ayırma", "4,631 ms", "İhmal edilebilir"],
            ["Kalite ölçümü", "17,961 ms", "İhmal edilebilir"],
            ["Agent ASR", "792,453 ms", "Gerçek zamanın belirgin önünde"],
            ["Customer ASR", "559,172 ms", "Gerçek zamanın belirgin önünde"],
            ["Timeline merge", "0,023 ms", "İhmal edilebilir"],
            ["Ollama analiz", "4.300,953 ms", "Toplam gecikmenin baskın bileşeni"],
        ],
        [3300, 1800, 4260],
    )
    doc.add_paragraph(
        "İki kanal ASR toplamı yaklaşık 1,352 sn'dir. 51,11 sn görüşme duvar "
        "süresine göre ASR RTF yaklaşık 0,026'dır. Model yükleme süresi aşama "
        "izlerine dahil değildir."
    )

    doc.add_heading("7. Timeline ve overlap düzeltmesi", level=1)
    add_callout(
        doc,
        "Doğrulama sonucu",
        "İlk uygulama geniş Whisper segmenti nedeniyle 4 overlap üretmişti. "
        "Otomatik word-gap bölme ve kanal bazlı VAD sonrasında beklenen 2 olayın "
        "2'si de bulundu; false positive kalmadı.",
        GREEN,
    )
    add_table(
        doc,
        ["Kaynak", "Overlap aralıkları / sayı"],
        [
            ["İnsan referansı", "29,569–29,810 sn ve 43,639–44,320 sn — toplam 2"],
            ["İlk uygulama", "Geniş customer segmenti nedeniyle toplam 4"],
            ["Güncel sistem", "29,63 sn ve 43,68 sn çevresinde toplam 2 — PASS"],
        ],
        [2700, 6660],
    )
    doc.add_paragraph(
        "Whisper customer kanalındaki 27,950–49,180 sn arasındaki ayrı kısa "
        "konuşmaları tek geniş segmentte döndürmüştü. TranscriptSegmentSplitter "
        "kelime zamanlarındaki 0,8 sn üzeri sessizliklerden bu çıktıyı otomatik "
        "böler. EnergySpeechActivityDetector her kanalın enerji aralıklarını "
        "bağımsız çıkarır; TimelineMerger overlap kararını metin segmentinin geniş "
        "sınırından değil iki kanalın gerçek konuşma aralığı kesişiminden verir."
    )
    doc.add_heading("7.1 Güncel tam konuşmanın kritik bölümü", level=2)
    doc.add_paragraph(
        "[27,950–29,840] customer: My debit card.\n"
        "[29,630–30,910] agent: Can you repeat that please?\n"
        "[34,320–36,540] customer: Yes. My debit card.\n"
        "[36,270–38,650] agent: Is there anything else I can help you with today?\n"
        "[42,780–44,680] customer: No. That was going to be it.\n"
        "[43,680–44,500] agent: Thank you for calling.\n"
        "[44,660–45,480] agent: Have a great day.\n"
        "[48,960–49,180] customer: Bye."
    )
    doc.add_paragraph(
        "Bu işlem referans metinden veya elle yapılan bir düzeltmeden yararlanmaz. "
        "reference.json yalnız sonuç üretildikten sonra test ve kabul ölçümü için kullanılır."
    )

    doc.add_heading("8. Analiz katmanı değerlendirmesi", level=1)
    add_table(
        doc,
        ["Alan", "Çıktı", "Değerlendirme"],
        [
            ["Need/topic", "debit card replacement", "Doğru"],
            ["Intent", "request", "Doğru"],
            ["Emotion", "neutral", "Makul; yalnız metne dayalı kesin duygu sayılmamalı"],
            ["Conversation stage", "initial request", "Tam görüşme için yanlış; closing/completed beklenir"],
            ["Recommended action", "replace debit card", "Ana ihtiyaçla uyumlu fakat süreç/politika doğrulaması yok"],
            ["Evidence", "8 segment", "Geçerli ID'ler; fakat gereğinden geniş ve agent kapanışlarını içeriyor"],
        ],
        [2100, 2500, 4760],
    )
    doc.add_paragraph(
        "Ollama adaptörü JSON ayrıştırma, zorunlu alan, confidence aralığı ve "
        "timeline dışı evidence ID üretimini engellemektedir. Semantik kalite için "
        "ek politika gereklidir: conversation_stage son event'lere göre kontrol "
        "edilmeli, evidence minimum ve konuya doğrudan bağlı olmalı, ürün/işlem "
        "önerileri dış katalog veya kural katmanıyla doğrulanmalıdır."
    )

    doc.add_heading("9. İzlenebilirlik ve yeniden üretilebilirlik", level=1)
    doc.add_paragraph(
        "Audit olay sırası: stereo.input.validated → channel.transcribed (agent) → "
        "channel.transcribed (customer) → timeline.concatenated_by_timestamp → "
        "analysis.request → analysis.response → validation.completed."
    )
    add_bullet(doc, "audit.json: kaynak yolu, aktif kanal, iki kanalın ayrı transcript'i, tam birleşik transcript, Ollama isteği/yanıtı ve validasyon.")
    add_bullet(doc, "result.json: kaynak, kanal haritası, kalite, segment/kelime zamanları, overlap, analiz ve aşama izleri.")
    add_bullet(doc, "report.md: insan okunur kalite, timeline, analiz ve süre özeti.")
    add_bullet(doc, "reference.json: insan transkripti, speaker_role, segment ve kelime zamanları.")
    add_bullet(
        doc,
        "Deterministik fixture üreticileri sentetik stereo ve iki mono kaydı ortak "
        "eksene yerleştirir; CLI'da tr/en/auto dil seçimi canlı .env'i değiştirmeden "
        "offline test sağlar ve model sınırları fake adaptörlerle doğrulanır.",
    )

    doc.add_heading("10. Riskler ve kapanış planı", level=1)
    add_table(
        doc,
        ["Öncelik", "İş", "Kapanış ölçütü"],
        [
            ["Tamamlandı", "Kanal bazlı VAD ve overlap düzeltmesi", "HarperValley referansında 2/2 overlap; false positive 0"],
            ["Tamamlandı", "Golden referans değerlendirici", "WER, speaker attribution, timing drift ve overlap otomatik raporlanır"],
            ["P1", "Analiz evresi/evidence politikası", "Full-call stage=closing/completed; evidence konuya bağlı minimum set"],
            ["P1", "Model benchmark", "tiny/small için WER, RTF ve bellek karşılaştırması"],
            ["P1", "İnsan dinleme kabulü", "Kanal, metin ve duyulan zamanlama bağımsız bir dinleyici tarafından imzalanır"],
            ["P2", "Crosstalk/dual-mono kalite bayrakları", "Eşikler etiketli fixture setiyle kalibre edilir"],
            ["P2", "Canlı entegrasyon tasarımı", "Offline kapanışından sonra ayrı faz olarak ele alınır"],
        ],
        [1100, 3500, 4760],
    )

    doc.add_heading("11. Sonuç", level=1)
    doc.add_paragraph(
        "Offline stereo mimari amaca uygun biçimde kurulmuş; mevcut canlı sistemden "
        "izole, sınıf tabanlı, test edilebilir ve raporlanabilir durumdadır. Kanal "
        "ayrımı, konuşmacı kimliği, örnek ekseni, transkripsiyon ve analiz zinciri "
        "gerçek veri üzerinde çalışmaktadır. Kanal bazlı VAD ve golden referans "
        "validasyonu tamamlanmış; gerçek çağrı overall_status=pass üretmiştir. "
        "Bununla birlikte ASR kusursuz değildir, Ollama'nın conversation_stage "
        "çıktısı semantik kabulü geçmemiştir ve otomatik ajan fiziksel dinleme "
        "onayı vermemiştir. Bu nedenle teknik zamanlama/kanal kabulü PASS, genel "
        "ürün kabulü ise belirtilen P1 maddelerine bağlı koşullu kabul olarak değerlendirilmelidir."
    )
    add_callout(
        doc,
        "Önerilen karar",
        "Offline görev kapsamı teknik olarak karşılanmıştır. Canlı entegrasyon ayrı "
        "fazdır; öncesinde small model benchmark'ı, analiz semantik politikası ve "
        "insan dinleme kabulü tamamlanmalıdır.",
        GREEN,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build().resolve())
